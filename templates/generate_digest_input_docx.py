# -*- coding: utf-8 -*-
"""Шаблон Word для сбора/обновления данных корпоративного дайджеста.

Палитра Virtu/JX: #C00000 / #1A1A1A (без синих акцентов).
Запуск: python templates/generate_digest_input_docx.py
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BRAND = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0xC0, 0x00, 0x00)
ACCENT_LIGHT = RGBColor(0xFB, 0x5A, 0x5A)
TEXT = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEADER_BG = "C00000"
ALT_ROW_BG = "FFF0F0"
META_BG = "F5F5F5"
HINT_BG = "FFF8F8"

FONT = "Calibri"

ROOT = Path(__file__).resolve().parents[1]
OUT_DIRS = [
    ROOT / "docs",
    ROOT / "output",
]
OUT_NAME = "шаблон_ввода_данных_дайджест.docx"


def shade(cell, fill_hex: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_run(p, text, size=11, bold=False, color=TEXT, italic=False):
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    return run


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(p, text, size=22, bold=True, color=BRAND)
    p.paragraph_format.space_after = Pt(4)


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    set_run(p, text, size=12, bold=True, color=ACCENT_LIGHT)
    p.paragraph_format.space_after = Pt(10)


def add_section(doc, text):
    p = doc.add_paragraph()
    set_run(p, text.upper(), size=12, bold=True, color=ACCENT)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "C00000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_hint(doc, text):
    p = doc.add_paragraph()
    set_run(p, text, size=10, italic=True, color=MUTED)
    p.paragraph_format.space_after = Pt(8)


def fill_cell(cell, text, bold=False, color=TEXT, size=10, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p, text, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        fill_cell(cell, h, bold=True, color=WHITE, size=10)
        shade(cell, HEADER_BG)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            fill_cell(cell, str(val), size=10)
            if i % 2 == 1:
                shade(cell, ALT_ROW_BG)
    if col_widths:
        for row in table.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = Cm(w)
    doc.add_paragraph()
    return table


def add_meta_table(doc, pairs):
    table = doc.add_table(rows=len(pairs), cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate(pairs):
        fill_cell(table.rows[i].cells[0], k, bold=True, size=10)
        fill_cell(table.rows[i].cells[1], v, size=10)
        shade(table.rows[i].cells[0], META_BG)
        shade(table.rows[i].cells[1], META_BG)
    table.rows[0].cells[0].width = Cm(5)
    table.rows[0].cells[1].width = Cm(12)
    doc.add_paragraph()


def blank_rows(n, cols):
    return [["…" for _ in range(cols)] for _ in range(n)]


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)
    style.font.color.rgb = TEXT

    add_title(doc, "Шаблон ввода данных дайджеста")
    add_subtitle(doc, "Virtu Systems · Корпоративный статус проектов")
    add_hint(
        doc,
        "Заполните таблицы и поля ниже. Сохраните файл и передайте в проектный офис / AI-помощнику "
        "для обновления HTML / PPTX / MD. Не меняйте названия разделов и заголовки колонок.",
    )

    add_section(doc, "1. Метаданные периода")
    add_meta_table(
        doc,
        [
            ("Период отчёта", "с января по … 202… г."),
            ("Месяц дайджеста", "например: май 2026"),
            ("Рук. тех. дирекции", "Столбов А."),
            ("Рук. проектного офиса", "Сорванов О."),
            ("Дата подготовки", "ДД.ММ.ГГГГ"),
            ("Версия / комментарий", "…"),
        ],
    )

    add_section(doc, "2. Портфель по страховым компаниям")
    add_hint(doc, "Одна строка = один проект. «В работе» и «Выполнено» — число или краткий текст.")
    add_table(
        doc,
        ["СК", "Проект", "Выполнено", "В работе"],
        [
            ["РГС", "VFOS", "…", "…"],
            ["РГС", "ОФР (ИФЛ, Базовая, Сервисная)", "…", "…"],
            ["РГС", "Аутстафф GuideWire / Импульс", "…", "—"],
            ["РГС", "РГС Ипотека", "…", "…"],
            ["Согласие", "Согласие Ипотека", "…", "…"],
            ["Согласие", "Вита / НПФ / Банки", "…", "…"],
            ["Согласие", "VFOS Согласие", "…", "—"],
            ["Альфа", "VFOS Альфа", "…", "—"],
            ["Райффайзен", "VFOS Райффайзен", "…", "—"],
            ["Ренессанс", "Ренессанс Ипотека", "…", "…"],
            ["Инсайт", "Инсайт Ипотека", "—", "—"],
            ["ВСК/МСГ", "ВСК Автострахование, ВСК ДМС, МСГ", "…", "…"],
            ["—", "Virtu Drive", "…", "—"],
        ],
        col_widths=[3.2, 7.5, 3.2, 2.5],
    )

    add_section(doc, "3. Команды (численность, чел.)")
    add_hint(doc, "Укажите роли и количество. Итог посчитается при сборке презентации.")
    add_table(
        doc,
        ["Команда", "Роль 1 / чел.", "Роль 2 / чел.", "Роль 3 / чел.", "Роль 4 / чел."],
        [
            ["VFOS", "Аналитики / …", "Разработчики / …", "QA / …", "—"],
            ["ОФР", "TypeScript / …", "QA / …", "БА / …", "—"],
            ["Аутстафф", "GoSu / …", "Аналитик / …", "TypeScript / …", "Java / …"],
            ["Ипотека", "Аналитики / …", "Разработчики / …", "Настройщики / …", "—"],
            ["Согласие Вита/НПФ", "Аналитики / …", "Разработчики / …", "Настройщики / …", "—"],
            ["Согласие Ипотека", "Аналитики / …", "Разработчики / …", "Настройщики / …", "—"],
            ["ВСК / МСГ", "Аналитики / …", "Разработчики / …", "—", "—"],
            ["Virtu Drive", "Аналитик / …", "Java / …", "—", "—"],
        ],
        col_widths=[3.5, 3.2, 3.5, 3.5, 2.8],
    )

    add_section(doc, "4. VFOS — ключевые активности")
    add_meta_table(
        doc,
        [
            ("PM", "Сорванов О.Н."),
            ("Аккаунт", "Черноус С."),
            ("Выполнено / в работе", "… / …"),
        ],
    )
    add_table(
        doc,
        ["СК", "Направление", "Реализованный функционал", "План на 1-й год"],
        blank_rows(6, 4)[:3]
        + [
            ["РГС", "…", "…", "… млн ₽"],
            ["Альфа", "Инфраструктура", "…", "—"],
            ["Райффайзен", "Архитектура", "…", "—"],
        ],
        col_widths=[2.2, 3.5, 8.0, 2.5],
    )

    add_section(doc, "5. ОФР и Аутстафф")
    add_meta_table(
        doc,
        [
            ("PM", "Сорванов О."),
            ("Аккаунт", "Черноус С."),
            ("ОФР, чел.", "…"),
            ("Аутстафф, чел.", "…"),
        ],
    )
    add_hint(doc, "Блоки: ОФР.ИФЛ · ОФР.Базовая/Сервисная · ОФР.ИФЛ и Базовая · Аутстафф · В работе")
    add_table(
        doc,
        ["Блок", "Активность (кратко)", "Статус (готово / в работе)"],
        blank_rows(8, 3),
        col_widths=[4.0, 9.5, 3.0],
    )

    add_section(doc, "6. Ипотечный портфель")
    add_meta_table(
        doc,
        [
            ("PM", "Кузнецова А."),
            ("Аккаунт", "Черноус С., Клейн Т., Конкин М."),
            ("Команда, чел.", "…"),
            ("Инциденты за месяц", "…"),
        ],
    )
    add_table(
        doc,
        ["Площадка", "Реализованный функционал", "Сделано / в работе"],
        [
            ["РГС Ипотека", "…", "… / …"],
            ["Ренессанс Ипотека", "…", "… / …"],
            ["Инсайт Ипотека", "…", "—"],
        ],
        col_widths=[3.5, 10.0, 3.0],
    )

    add_section(doc, "7. Согласие / ВСК / МСГ")
    add_meta_table(
        doc,
        [
            ("PM", "Кочетков Н., Царёва А."),
            ("Аккаунт", "Клейн Т., Конкин М."),
            ("Согласие Ипотека вып/в раб.", "… / …"),
            ("Вита/НПФ/Банки/… вып/в раб.", "… / …"),
        ],
    )
    add_table(
        doc,
        ["Блок", "Детализация", "Результат"],
        [
            ["Интеграция ВТБ ELM", "…", "…"],
            ["Спецусловия партнёров", "…", "—"],
            ["Журнал пролонгации", "…", "…"],
            ["Вита / НПФ / Банки", "…", "… вып / … в работе"],
            ["МП Эльбрус", "…", "…"],
            ["ВСК / МСГ", "…", "—"],
        ],
        col_widths=[3.8, 9.2, 3.5],
    )

    add_section(doc, "8. Техническая поддержка")
    add_meta_table(
        doc,
        [
            ("Рук. поддержки", "Смирнова А."),
            ("Рук. тех. дирекции", "Столбов А."),
            ("Период ТП", "с января по …"),
        ],
    )
    add_table(
        doc,
        ["Система / блок", "Категория 1", "Категория 2", "Категория 3", "Категория 4", "Всего"],
        [
            ["VFOS", "Доступы …", "Изм. БД …", "Инцид. …", "Прод./Проч. …", "…"],
            ["Согласие Ипотека", "Конс. …", "Нетип. …", "—", "—", "…"],
            ["Ипотека (РГС, Рен., Инс.)", "Нетип. …", "Подкл. …", "Конс. …", "Аварии …", "…"],
            ["Альфа VFOS", "…", "…", "…", "…", "…"],
            ["Virtu Drive", "Конс. …", "Нетип. …", "—", "—", "…"],
            ["Согласие VFOS", "…", "…", "…", "…", "…"],
        ],
        col_widths=[3.5, 2.5, 2.5, 2.5, 2.8, 1.8],
    )

    add_section(doc, "9. VirtuDrive / Pre-sale")
    add_meta_table(doc, [("PM", "Шейко Н."), ("Аккаунт", "Чесноков Д.")])
    add_table(
        doc,
        ["Проект", "Статус", "Следующий шаг", "Дедлайн"],
        [
            ["Virtu Drive", "…", "…", "—"],
            ["Тендер Автофинанс — Модуль Страховки", "…", "…", "ДД.ММ.ГГГГ"],
            ["Мобильное приложение Агента ЛК", "…", "…", "ДД.ММ.ГГГГ"],
            ["Конструктор комиссионных продуктов", "…", "…", "—"],
        ],
        col_widths=[5.5, 3.0, 5.0, 2.8],
    )

    add_section(doc, "10. Благодарность сотрудникам")
    add_table(
        doc,
        ["Площадка", "Сотрудники (ФИО)", "Реализованный функционал"],
        blank_rows(10, 3),
        col_widths=[3.5, 5.5, 7.5],
    )

    add_section(doc, "11. Матрица ответственности PM")
    add_table(
        doc,
        ["PM", "Проекты", "Разработка (вып. / в раб.)", "ТП (задач)"],
        [
            ["Сорванов О.Н.", "…", "… / …", "…"],
            ["Кузнецова А.", "…", "… / …", "…"],
            ["Кочетков Н., Царёва А.", "…", "… / …", "—"],
            ["Кочетков Н., Царёва А.", "Согласие Ипотека", "… / …", "…"],
            ["Кочетков Н.", "ВСК …", "… / …", "—"],
            ["Царёва А.", "МП Эльбрус", "—", "—"],
            ["Шейко Н.", "Virtu Drive, …", "…", "…"],
        ],
        col_widths=[3.8, 6.5, 3.5, 3.0],
    )

    add_section(doc, "12. Свободные комментарии / риски / фокус следующего месяца")
    add_hint(doc, "Любой текст: блокеры, договорённости, что показать на статусе.")
    for _ in range(4):
        p = doc.add_paragraph()
        set_run(p, "• …", size=11, color=TEXT)

    add_section(doc, "Инструкция для сборки")
    for line in [
        "1. Заполните разделы 1–12; пустые строки «…» замените фактами или удалите.",
        "2. Сохраните файл как дайджест_<месяц>_<год>_ввод.docx.",
        "3. Передайте файл / приложите в чат Cursor — по шаблону обновят HTML, MD и PPTX.",
        "4. Не меняйте заголовки разделов — по ним парсится структура.",
        "5. Числа: разделитель тысяч — пробел (3 897); доли — «вып / в раб.».",
    ]:
        p = doc.add_paragraph()
        set_run(p, line, size=10, color=TEXT)

    paths = []
    for d in OUT_DIRS:
        d.mkdir(parents=True, exist_ok=True)
        path = d / OUT_NAME
        doc.save(path)
        paths.append(path)
    return paths


if __name__ == "__main__":
    for p in build():
        print(f"Saved: {p}")
